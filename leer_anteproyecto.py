# -*- coding: utf-8 -*-
"""Extraer TODO el contenido del anteproyecto Word: párrafos, tablas completas."""
import sys
from docx import Document

RUTA = r"c:\Users\gusta\OneDrive - UNIR\TRABAJOS DE GRADO\TRABAJOS DE GRADO\UAN\LUIS\01_PARA_PRESENTAR\Documento Anteproyecto_Mayo_7.docx"
doc = Document(RUTA)

print("=" * 90)
print(f"EXTRACCIÓN ANTEPROYECTO: {RUTA}")
print("=" * 90)

print("\n# 1. PÁRRAFOS (títulos y texto):")
print("-" * 90)
for i, para in enumerate(doc.paragraphs, 1):
    texto = para.text.strip()
    if not texto:
        continue
    # Obtener estilo del párrafo para detectar encabezados
    estilo = para.style.name if para.style else "Normal"
    prefijo = ""
    if "Heading 1" in estilo or "Título 1" in estilo: prefijo = "# "
    elif "Heading 2" in estilo or "Título 2" in estilo: prefijo = "## "
    elif "Heading 3" in estilo or "Título 3" in estilo: prefijo = "### "
    elif "Heading 4" in estilo or "Título 4" in estilo: prefijo = "#### "
    elif "Title" in estilo or "Título" in estilo: prefijo = "# DOCUMENTO: "
    elif "Subtitle" in estilo: prefijo = ">> "
    print(f"{prefijo}[{i:03d}|{estilo}]  {texto}")

print("\n\n# 2. TABLAS DEL DOCUMENTO (Cronograma, Objetivos, Módulos, etc.):")
print("-" * 90)
for idx_t, tabla in enumerate(doc.tables, 1):
    print(f"\n==================== TABLA #{idx_t} (filas={len(tabla.rows)}, cols={len(tabla.columns)}) ====================")
    for idx_f, fila in enumerate(tabla.rows, 1):
        celdas = [celda.text.strip().replace("\n", " | ") for celda in fila.cells]
        print(f"F{idx_f:02d}: " + " || ".join(celdas))

print("\n\nFIN EXTRACCIÓN")
